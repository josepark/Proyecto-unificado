"""
Comando de importacion: carga la Matriz_Activos_identificados.xlsx
al modelo relacional del inventario SGSI.

Uso:
    python manage.py importar_matriz --file data/Matriz_Activos_identificados.xlsx
"""
import re
from datetime import datetime

from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from inventario.models import (
    Activo, ActivoInfraestructura, SistemaInformacion,
    Zona, VLAN, AmenazaMITRE, ControlISO)

# Diccionario de descripciones legibles para tacticas/tecnicas MITRE frecuentes
MITRE_DESC = {
    "T1486": "Data Encrypted for Impact (Ransomware)",
    "T1190": "Exploit Public-Facing Application",
    "T1611": "Escape to Host (Escape de contenedor)",
    "T1078": "Valid Accounts",
    "T1485": "Data Destruction",
    "T1490": "Inhibit System Recovery",
    "T1530": "Data from Cloud Storage / NAS",
    "T1005": "Data from Local System",
    "T1213": "Data from Information Repositories",
    "T1499": "Endpoint Denial of Service",
    "T1498": "Network Denial of Service (DDoS)",
    "T1008": "Fallback Channels",
    "T1557": "Adversary-in-the-Middle",
    "T1040": "Network Sniffing",
    "T1200": "Hardware Additions",
    "T1133": "External Remote Services",
    "T1110": "Brute Force",
    "TA0040": "Impact",
    "TA0005": "Defense Evasion",
    "TA0001": "Initial Access",
    "TA0008": "Lateral Movement",
    "TA0009": "Collection",
}

CLASIF_MAP = {
    "altamente confidencial": Activo.Clasificacion.ALTAMENTE,
    "confidencial": Activo.Clasificacion.CONFIDENCIAL,
    "uso interno": Activo.Clasificacion.INTERNO,
    "publico": Activo.Clasificacion.PUBLICO,
    "público": Activo.Clasificacion.PUBLICO,
}

RIESGO_MAP = {
    "critico": Activo.NivelRiesgo.CRITICO, "crítico": Activo.NivelRiesgo.CRITICO,
    "alto": Activo.NivelRiesgo.ALTO, "medio": Activo.NivelRiesgo.MEDIO,
    "bajo": Activo.NivelRiesgo.BAJO,
}


def cid(val):
    """Convierte '4 - Critico' -> 4."""
    if val is None:
        return None
    m = re.match(r"\s*(\d)", str(val))
    return int(m.group(1)) if m else None


def norm(v):
    return (str(v).strip() if v is not None else "")


def parse_fecha(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    return None


def split_multi(texto):
    """Divide 'T1486 Ransomware / T1190 Exploit' en items."""
    if not texto:
        return []
    return [p.strip() for p in re.split(r"[/;\n]", str(texto)) if p.strip()]


class Command(BaseCommand):
    help = "Importa la matriz de activos SUIIN-SGSI-INV-001 al modelo relacional"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True)

    def get_amenaza(self, item):
        m = re.match(r"(T\d{4}|TA\d{4})", item)
        if not m:
            return None
        codigo = m.group(1)
        desc_resto = item[len(codigo):].strip()
        desc = MITRE_DESC.get(codigo, desc_resto[:200])
        obj, _ = AmenazaMITRE.objects.get_or_create(
            codigo=codigo, defaults={"descripcion": desc})
        if not obj.descripcion and desc:
            obj.descripcion = desc
            obj.save()
        return obj

    def get_control(self, item):
        m = re.match(r"(POL-SI-\d+|MAN-POL-SI-\d+|\d+\.\d+)", item)
        if not m:
            return None
        codigo = m.group(1)
        desc = item[len(codigo):].strip()[:200]
        obj, created = ControlISO.objects.get_or_create(
            codigo=codigo, defaults={"descripcion": desc})
        if created and desc:
            obj.descripcion = desc
            obj.save()
        return obj

    @transaction.atomic
    def handle(self, *args, **opts):
        wb = load_workbook(opts["file"], data_only=True)
        n_infra = 0
        n_sist = 0

        # ---- Hojas de infraestructura ----
        for sheet in ["Infraestructura de Red Servidor",
                      "Infraestructura de Red Swith y "]:
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=3, values_only=True):
                if not row or not row[0] or not str(row[0]).startswith("RED-"):
                    continue
                (id_activo, nombre, descripcion, tipo, subtipo, zona_txt,
                 ip_seg, vlan_txt, modelo, serial, clasif_txt, c, i, d,
                 valor, riesgo_txt, amenazas_txt, controles_txt, estado_txt,
                 fecha, notas) = row[:21]

                activo, _ = Activo.objects.update_or_create(
                    id_activo=norm(id_activo),
                    defaults=dict(
                        nombre=norm(nombre),
                        descripcion=norm(descripcion),
                        clase="INFRA",
                        clasificacion_si=CLASIF_MAP.get(norm(clasif_txt).lower(), ""),
                        confidencialidad=cid(c),
                        integridad=cid(i),
                        disponibilidad=cid(d),
                        nivel_riesgo=RIESGO_MAP.get(norm(riesgo_txt).lower(),
                                                    Activo.NivelRiesgo.SIN),
                        estado=Activo.Estado.ACTIVO,
                        fecha_registro=parse_fecha(fecha),
                        notas_seguridad=norm(notas),
                    ))

                zona = None
                if norm(zona_txt):
                    zona, _ = Zona.objects.get_or_create(nombre=norm(zona_txt))
                vlan = None
                if norm(vlan_txt) and norm(vlan_txt) != "—":
                    vlan, _ = VLAN.objects.get_or_create(etiqueta=norm(vlan_txt))

                ActivoInfraestructura.objects.update_or_create(
                    activo=activo,
                    defaults=dict(tipo=norm(tipo), subtipo=norm(subtipo),
                                  zona=zona, ip_segmento=norm(ip_seg), vlan=vlan,
                                  modelo=norm(modelo), serial_placa=norm(serial)))

                for it in split_multi(amenazas_txt):
                    a = self.get_amenaza(it)
                    if a:
                        activo.amenazas.add(a)
                for it in split_multi(controles_txt):
                    ctrl = self.get_control(it)
                    if ctrl:
                        activo.controles.add(ctrl)
                n_infra += 1

        # ---- Hoja de sistemas de informacion ----
        ws = wb["Sistemas de Información"]
        rows = list(ws.iter_rows(min_row=4, max_row=22, values_only=True))
        estado_op_map = {
            "en operación": SistemaInformacion.EstadoOperativo.OPERACION,
            "en operacion": SistemaInformacion.EstadoOperativo.OPERACION,
            "en implementación": SistemaInformacion.EstadoOperativo.IMPLEMENTACION,
            "en implementacion": SistemaInformacion.EstadoOperativo.IMPLEMENTACION,
            "no operación": SistemaInformacion.EstadoOperativo.NO_OPERACION,
            "no operacion": SistemaInformacion.EstadoOperativo.NO_OPERACION,
        }
        for row in rows:
            num = row[0]
            modulo = norm(row[1])
            if not modulo:
                continue
            n = int(float(num)) if num is not None else n_sist + 1
            id_activo = f"SIS-{n:03d}"
            estado_op_txt = norm(row[2]).lower()
            estado_op = estado_op_map.get(estado_op_txt,
                                          SistemaInformacion.EstadoOperativo.SIN_DATO)

            estado_activo = Activo.Estado.ACTIVO
            if estado_op == SistemaInformacion.EstadoOperativo.NO_OPERACION:
                estado_activo = Activo.Estado.INACTIVO
            elif estado_op == SistemaInformacion.EstadoOperativo.IMPLEMENTACION:
                estado_activo = Activo.Estado.IMPLEMENTACION

            activo, _ = Activo.objects.update_or_create(
                id_activo=id_activo,
                defaults=dict(
                    nombre=modulo,
                    descripcion=f"Modulo/sistema de informacion SUIIN: {modulo}.",
                    clase="SIST",
                    clasificacion_si=CLASIF_MAP.get(norm(row[10]).lower(), ""),
                    estado=estado_activo,
                ))

            prioridad_txt = norm(row[15]).upper()
            prioridad = (SistemaInformacion.Prioridad.ALTA
                         if "PRIORIZAR" in prioridad_txt
                         else SistemaInformacion.Prioridad.SIN)

            api_rest = None
            if norm(row[6]).lower() == "sí" or norm(row[6]).lower() == "si":
                api_rest = True

            sis, _ = SistemaInformacion.objects.update_or_create(
                activo=activo,
                defaults=dict(
                    estado_operativo=estado_op,
                    backend=norm(row[3]), frontend=norm(row[4]),
                    schema_bd=norm(row[5]), api_rest_nativa=api_rest,
                    integracion_gateway=norm(row[7]),
                    estado_documentacion=norm(row[8]),
                    sistema_mca_equivalente=norm(row[9]),
                    servidor_virtual=norm(row[13]).replace(".0", ""),
                    url=norm(row[14]),
                    priorizar_analisis=prioridad,
                ))

            # Accesos por rol: deprecado (Ola 4) — la Matriz RBAC es la fuente canónica.
            # La columna de roles MCA del Excel ya no se importa al Inventario.

            n_sist += 1

        # ---- Seccion 2: requisitos API Gateway ----
        gw_rows = list(ws.iter_rows(min_row=26, max_row=37, values_only=True))
        def chk(v):
            return norm(v) == "✓"
        for row in gw_rows:
            modulo = norm(row[0])
            if not modulo:
                continue
            sis = SistemaInformacion.objects.filter(
                activo__nombre__iexact=modulo).first()
            if not sis:
                continue
            sis.gw_validacion_jwt = chk(row[2])
            sis.gw_sso = chk(row[3])
            sis.gw_cors = chk(row[4])
            sis.gw_inyeccion_roles = chk(row[5])
            sis.gw_refresh_token = chk(row[6])
            sis.gw_balanceo_carga = chk(row[7])
            sis.save()

        self.stdout.write(self.style.SUCCESS(
            f"Importacion completa: {n_infra} activos de infraestructura, "
            f"{n_sist} sistemas de informacion."))
        self.stdout.write(
            f"Amenazas MITRE: {AmenazaMITRE.objects.count()} | "
            f"Controles ISO/POL: {ControlISO.objects.count()}")

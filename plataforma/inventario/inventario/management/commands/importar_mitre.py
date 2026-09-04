"""
Importa el catalogo oficial MITRE ATT&CK (enterprise) al modelo AmenazaMITRE.
Carga tacticas, tecnicas y subtecnicas con nombre, descripcion, tacticas
asociadas, plataformas y URL de referencia.

Uso:
    python manage.py importar_mitre --file data/enterprise-attack-v19_1.xlsx
"""
from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from inventario.models import AmenazaMITRE


def norm(v):
    return (str(v).strip() if v is not None else "")


class Command(BaseCommand):
    help = "Importa el catalogo MITRE ATT&CK enterprise al catalogo de amenazas"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True)

    @transaction.atomic
    def handle(self, *args, **opts):
        wb = load_workbook(opts["file"], read_only=True, data_only=True)
        n_ta = n_te = n_st = 0

        # --- Tacticas ---
        ws = wb["tactics"]
        rows = ws.iter_rows(min_row=2, values_only=True)
        for r in rows:
            cod = norm(r[0])
            if not cod:
                continue
            AmenazaMITRE.objects.update_or_create(codigo=cod, defaults=dict(
                nombre=norm(r[2]), descripcion=norm(r[3])[:2000],
                tipo=AmenazaMITRE.Tipo.TACTICA, url=norm(r[4]), version=norm(r[8])))
            n_ta += 1

        # --- Tecnicas y subtecnicas ---
        ws = wb["techniques"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            cod = norm(r[0])
            if not cod:
                continue
            es_sub = norm(r[11]).lower() in ("true", "1", "yes", "si")
            AmenazaMITRE.objects.update_or_create(codigo=cod, defaults=dict(
                nombre=norm(r[2]), descripcion=norm(r[3])[:4000],
                tipo=AmenazaMITRE.Tipo.SUBTECNICA if es_sub else AmenazaMITRE.Tipo.TECNICA,
                tacticas=norm(r[9])[:300], plataformas=norm(r[10])[:300],
                codigo_padre=norm(r[12]), url=norm(r[4]), version=norm(r[8])))
            if es_sub:
                n_st += 1
            else:
                n_te += 1

        total = AmenazaMITRE.objects.count()
        self.stdout.write(self.style.SUCCESS(
            f"Catalogo MITRE importado: {n_ta} tacticas, {n_te} tecnicas, "
            f"{n_st} subtecnicas. Total en catalogo de amenazas: {total}."))

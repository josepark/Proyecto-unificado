"""
Importación masiva de activos desde Excel — antes, registrar un lote
grande de equipos nuevos (tras una compra, por ejemplo) significaba
cargarlos uno por uno desde el formulario. Sigue el mismo patrón de dos
pasos que ya usaba RBAC para importar la matriz (rbac/negocio.py):
primero se analiza el
archivo sin guardar nada, se muestra fila por fila qué se va a crear y
qué tiene errores, y solo se escribe en la base de datos cuando el
usuario confirma explícitamente.

Reutiliza ActivoWriteSerializer para la validación fila por fila — la
misma que ya usa el formulario individual — en vez de reinventar reglas
de validación paralelas que podrían desalinearse con el tiempo.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

COLUMNAS = [
    "id_activo", "nombre", "clase", "clasificacion_si", "nivel_riesgo",
    "confidencialidad", "integridad", "disponibilidad", "estado",
    "ciclo_vida", "datacenter", "propietario", "custodio",
    "area_responsable", "procesa_datos_personales", "rto", "rpo",
    "descripcion",
]
# Encabezados legibles para la plantilla descargable — el análisis del
# archivo subido acepta cualquiera de los dos (clave técnica o este
# encabezado), sin distinguir mayúsculas/acentos, para no obligar a la
# persona a escribir exactamente "clasificacion_si".
ENCABEZADOS = {
    "id_activo": "ID Activo (vacío = automático)", "nombre": "Nombre *",
    "clase": "Clase * (INFRA/SIST/EQUI)",
    "clasificacion_si": "Clasificación SI (ALTA/CONF/INT/PUB)",
    "nivel_riesgo": "Nivel de riesgo (SIN/BAJO/MED/ALTO/CRIT)",
    "confidencialidad": "Confidencialidad (0-4)", "integridad": "Integridad (0-4)",
    "disponibilidad": "Disponibilidad (0-4)",
    "estado": "Estado (ACT/IMP/INA/RET)",
    "ciclo_vida": "Ciclo de vida (PROD/PLAN/MANT/RETI)",
    "datacenter": "Centro de datos (código, ej. DC-POPAYAN)",
    "propietario": "Propietario", "custodio": "Custodio",
    "area_responsable": "Área responsable",
    "procesa_datos_personales": "Procesa datos personales (Si/No)",
    "rto": "RTO", "rpo": "RPO", "descripcion": "Descripción",
}
_CAMPOS_ENTEROS = ("confidencialidad", "integridad", "disponibilidad")
_FILA_EJEMPLO = {
    "id_activo": "", "nombre": "Switch de borde — Sala de Juntas",
    "clase": "INFRA", "clasificacion_si": "INT", "nivel_riesgo": "MED",
    "confidencialidad": 2, "integridad": 2, "disponibilidad": 2,
    "estado": "ACT", "ciclo_vida": "PROD", "datacenter": "DC-POPAYAN",
    "propietario": "Coordinación TI", "custodio": "", "area_responsable": "",
    "procesa_datos_personales": "No", "rto": "4 horas", "rpo": "1 hora",
    "descripcion": "",
}


def generar_plantilla():
    """Devuelve un Workbook con encabezados, una fila de ejemplo y una
    segunda hoja con los valores válidos de cada campo de opciones."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Activos a importar"
    ws.append([ENCABEZADOS[c] for c in COLUMNAS])
    fill = PatternFill("solid", fgColor="0d2b23")
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = fill
    ws.append([_FILA_EJEMPLO[c] for c in COLUMNAS])
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(w, 45)

    ref = wb.create_sheet("Valores válidos")
    ref.append(["Campo", "Valores aceptados"])
    from .models import Activo, ClaseActivo
    codigos = list(ClaseActivo.objects.filter(activo=True).order_by("orden", "codigo")
                   .values_list("codigo", flat=True))
    if not codigos:
        codigos = [c[0] for c in Activo.Clase.choices]
    ref.append(["clase", " / ".join(codigos)])
    ref.append(["clasificacion_si", " / ".join(c[0] for c in Activo.Clasificacion.choices)])
    ref.append(["nivel_riesgo", " / ".join(c[0] for c in Activo.NivelRiesgo.choices)])
    ref.append(["estado", " / ".join(c[0] for c in Activo.Estado.choices)])
    ref.append(["ciclo_vida", " / ".join(c[0] for c in Activo.CicloVida.choices)])
    ref.append(["confidencialidad/integridad/disponibilidad", "0 a 4"])
    ref.append(["procesa_datos_personales", "Si / No"])
    for col in ref.columns:
        w = max(len(str(c.value or "")) for c in col) + 2
        ref.column_dimensions[col[0].column_letter].width = min(w, 70)
    return wb


def _clave_normalizada(texto):
    import unicodedata
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return t.strip().lower().replace(" ", "_")


def _mapear_encabezados(fila_encabezados):
    """Acepta tanto la clave técnica ('clasificacion_si') como el
    encabezado legible de la plantilla ('Clasificación SI (...)') —
    normaliza y compara contra ambos."""
    inversa = {_clave_normalizada(v): k for k, v in ENCABEZADOS.items()}
    mapa = {}
    for i, h in enumerate(fila_encabezados):
        if h is None:
            continue
        norm = _clave_normalizada(h)
        if norm in COLUMNAS:
            mapa[i] = norm
        elif norm in inversa:
            mapa[i] = inversa[norm]
        else:
            # el encabezado legible tiene texto extra despues del campo
            # (ej. "clase_*_(infra/sist/equi)") — probar por coincidencia
            # de prefijo con cada clave conocida.
            for clave in COLUMNAS:
                if norm.startswith(clave):
                    mapa[i] = clave
                    break
    return mapa


def _completar_bloque_detalle(cuerpo):
    """La plantilla Excel no trae bloques anidados (infra/sistema/equipo);
    los completa vacíos según el catálogo ClaseActivo para pasar la misma
    validación que el formulario web."""
    from .models import ClaseActivo

    clase = (cuerpo.get("clase") or "").strip()
    if not clase:
        return cuerpo
    cat = ClaseActivo.objects.filter(codigo=clase, activo=True).first()
    if not cat or cat.modelo_detalle in ("generico", "ninguno"):
        return cuerpo
    if cat.modelo_detalle not in cuerpo or cuerpo[cat.modelo_detalle] is None:
        cuerpo = dict(cuerpo)
        cuerpo[cat.modelo_detalle] = {}
    return cuerpo


def _normalizar_valor(campo, valor):
    if valor is None:
        return "" if campo not in _CAMPOS_ENTEROS else None
    if campo in _CAMPOS_ENTEROS:
        if valor == "":
            return None
        try:
            n = int(valor)
        except (TypeError, ValueError):
            raise ValueError(f"'{campo}' debe ser un número entero (0-4), recibido: {valor!r}")
        return n
    if campo == "procesa_datos_personales":
        return str(valor).strip().lower() in ("si", "sí", "true", "1", "yes")
    return str(valor).strip()


def analizar_archivo(archivo_django):
    """Lee el .xlsx subido y devuelve una lista de filas con su
    resultado de validación, sin guardar nada en la base de datos."""
    from .models import Datacenter
    from .serializers import ActivoWriteSerializer

    wb = load_workbook(archivo_django, read_only=True, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    mapa = _mapear_encabezados(filas[0])
    if "nombre" not in mapa.values() or "clase" not in mapa.values():
        raise ValueError(
            "El archivo no tiene las columnas mínimas requeridas (Nombre y Clase). "
            "Descargue la plantilla y no cambie los encabezados de la primera fila.")

    datacenters_por_codigo = {d.codigo.strip().lower(): d.id
                              for d in Datacenter.objects.all()}
    ids_ya_usados_en_lote = set()
    resultados = []

    for num_fila, fila in enumerate(filas[1:], start=2):
        if fila is None or not any(v not in (None, "") for v in fila):
            continue  # fila en blanco — se ignora, no cuenta como error
        crudo = {mapa[i]: v for i, v in enumerate(fila) if i in mapa}

        try:
            cuerpo = {c: _normalizar_valor(c, crudo.get(c)) for c in COLUMNAS}
        except ValueError as e:
            resultados.append({"fila": num_fila, "estado": "error",
                               "mensaje": str(e), "datos": crudo})
            continue

        dc_codigo = (cuerpo.pop("datacenter", "") or "").strip()
        if dc_codigo:
            dc_id = datacenters_por_codigo.get(dc_codigo.lower())
            if dc_id is None:
                resultados.append({"fila": num_fila, "estado": "error",
                                   "mensaje": f"Centro de datos '{dc_codigo}' no existe.",
                                   "datos": cuerpo})
                continue
            cuerpo["datacenter"] = dc_id

        id_activo = (cuerpo.get("id_activo") or "").strip()
        if id_activo:
            if id_activo in ids_ya_usados_en_lote:
                resultados.append({"fila": num_fila, "estado": "error",
                                   "mensaje": f"ID Activo '{id_activo}' repetido en este mismo archivo.",
                                   "datos": cuerpo})
                continue
            ids_ya_usados_en_lote.add(id_activo)
        else:
            cuerpo.pop("id_activo", None)

        cuerpo = _completar_bloque_detalle(cuerpo)
        ser = ActivoWriteSerializer(data=cuerpo)
        if ser.is_valid():
            resultados.append({"fila": num_fila, "estado": "ok",
                               "mensaje": "Listo para crear", "datos": cuerpo})
        else:
            mensaje = "; ".join(f"{campo}: {errs[0]}" for campo, errs in ser.errors.items())
            resultados.append({"fila": num_fila, "estado": "error",
                               "mensaje": mensaje, "datos": cuerpo})

    return resultados


def aplicar_filas(filas, usuario):
    """Vuelve a validar y crea cada fila marcada como lista para crear.
    Se re-valida (no se confía ciegamente en lo que ya analizó el
    navegador) porque pudo haber pasado tiempo entre analizar y
    confirmar — otra persona pudo haber tomado el mismo ID Activo, por
    ejemplo. Cada fila se crea de forma independiente: si una falla, no
    bloquea a las demás."""
    from .serializers import ActivoWriteSerializer

    creados, fallidos = [], []
    for f in filas:
        datos = _completar_bloque_detalle(dict(f.get("datos", {})))
        ser = ActivoWriteSerializer(data=datos)
        if ser.is_valid():
            activo = ser.save()
            creados.append({"fila": f.get("fila"), "id": activo.id,
                            "id_activo": activo.id_activo})
        else:
            mensaje = "; ".join(f"{campo}: {errs[0]}" for campo, errs in ser.errors.items())
            fallidos.append({"fila": f.get("fila"), "mensaje": mensaje})
    return creados, fallidos
